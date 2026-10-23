"""Employer-facing cover letters must look professional in both PDF and Word.

These checks exercise real renderer artifacts, including page fields, layout,
editable body text, and an unsplit signature for long applications.
"""

from __future__ import annotations

import pdfplumber
import pytest
from docx import Document

from utils.document_design import render_cover_letter_docx, render_cover_letter_pdf


def _letter_model(repeat: int = 1) -> dict:
    paragraph = (
        "My verified professional experience includes coordination of health "
        "programmes, reporting, clinical support and team supervision. "
        "I would welcome the opportunity to contribute these established skills."
    )
    return {
        "name": "Amina Example",
        "headline": "Application Letter",
        "contact": {
            "location": "Kabul, Afghanistan",
            "phone": "+93 700 555 111",
            "email": "amina@example.org",
        },
        "subject": "Vacancy HR-2026-012",
        "target_role": "Medical Officer",
        "target_org": "Example Health NGO",
        "reference": "HR-2026-012",
        "body_lines": ["Dear Hiring Committee,", *([paragraph] * repeat)],
        "signature": {
            "name": "Amina Example",
            "title": "Medical Doctor",
            "email": "amina@example.org",
            "phone": "+93 700 555 111",
        },
    }


def test_cover_letter_word_is_real_editable_a4_with_page_numbers(tmp_path):
    path = tmp_path / "modern-letter.docx"
    render_cover_letter_docx(_letter_model(), path)
    doc = Document(path)
    sec = doc.sections[0]
    assert sec.page_width.mm == pytest.approx(210, abs=.1)
    assert sec.page_height.mm == pytest.approx(297, abs=.1)
    assert sec.left_margin.mm == pytest.approx(20, abs=.1)
    assert sec.right_margin.mm == pytest.approx(20, abs=.1)
    assert doc.core_properties.title == "Amina Example — Cover Letter"
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "My verified professional experience" in text
    assert text.count("Sincerely,") == 1
    assert text.count("Vacancy HR-2026-012") == 1
    assert not doc.element.xpath(".//w:txbxContent | .//w:drawing")
    footer = sec.footer.paragraphs[0]
    assert footer._p.xpath(".//w:instrText")[0].text == "PAGE"
    assert "Page " in footer.text
    assert doc.paragraphs[0].style.name == "JF Name"


def test_word_signature_paragraphs_stick_together(tmp_path):
    path = tmp_path / "signature.docx"
    render_cover_letter_docx(_letter_model(repeat=15), path)
    doc = Document(path)
    signoff_index = next(i for i, p in enumerate(doc.paragraphs) if p.text == "Sincerely,")
    signature = doc.paragraphs[signoff_index:]
    assert [p.text for p in signature] == [
        "Sincerely,", "Amina Example", "Medical Doctor",
        "Email: amina@example.org", "Phone: +93 700 555 111",
    ]
    assert all(p.paragraph_format.keep_with_next is True for p in signature[:-1])
    assert signature[-1].paragraph_format.keep_with_next is False


def test_long_pdf_letter_keeps_signature_on_one_page(tmp_path):
    path = tmp_path / "letter-with-many-paragraphs.pdf"
    render_cover_letter_pdf(_letter_model(repeat=48), path)
    with pdfplumber.open(path) as pdf:
        assert len(pdf.pages) >= 2
        bodies = [p.extract_text() or "" for p in pdf.pages]
        signoff_pages = [i for i, body in enumerate(bodies) if "Sincerely," in body]
        assert len(signoff_pages) == 1
        last = bodies[signoff_pages[0]]
        assert "Amina Example" in last
        assert "Medical Doctor" in last
        assert "Email: amina@example.org" in last
        assert "Phone: +93 700 555 111" in last
        assert signoff_pages[0] == len(pdf.pages) - 1
