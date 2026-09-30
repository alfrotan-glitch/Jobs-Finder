from datetime import date
from pathlib import Path

from utils.discovery import Job, deduplicate_jobs
from utils.documents import generate_application_package, generate_tailored_documents, prepare_application_bundle
from utils.medical_matcher import match_job_against_profile
from utils.resume_parser import extract_resume_text, parse_cv_evidence
from utils import tracker


def test_text_cv_parsing_and_evidence(tmp_path):
    cv = tmp_path / "cv.txt"
    cv.write_text("MD physician with 5 years clinical experience. Languages: Dari, Pashto, English. HMIS and BPHS reporting.")
    text = extract_resume_text(str(cv))
    assert "5 years" in text
    evidence = parse_cv_evidence({"personal": {}}, str(cv))
    assert "md_degree" in evidence
    assert "clinical_experience_years" in evidence
    assert "language_dari" in evidence


def test_deduplicate_jobs_by_reference_and_url():
    a = Job("1", "Medical Officer", "NGO", "Kabul", "https://x/jobs/1", "https://x/apply/1", "test", metadata={"reference_number": "REF-1"})
    b = Job("2", "Medical Officer", "NGO", "Kabul", "https://x/jobs/duplicate", "https://x/apply/2", "test", metadata={"reference_number": "REF-1"})
    c = Job("3", "Medical Officer", "NGO", "Kabul", "https://x/jobs/3?utm=1", "https://x/jobs/3", "test")
    d = Job("4", "Medical Officer", "NGO", "Kabul", "https://x/jobs/3", "https://x/jobs/3", "test")
    unique = deduplicate_jobs([a, b, c, d])
    assert [j.id for j in unique] == ["1", "3"]


def test_tailored_documents_do_not_claim_missing_license():
    job = {
        "id": "j1",
        "title": "Medical Officer",
        "company": "NGO",
        "location": "Kabul",
        "url": "https://example.org/job",
        "apply_url": "https://example.org/apply",
        "description": "Medical Officer. MD and valid medical license required. Closing date: 30 September 2026.",
        "metadata": {},
    }
    profile = {
        "personal": {"first_name": "A", "last_name": "Doctor", "email": "a@example.org", "location": "Kabul"},
        "medical_education": [{"degree": "MD", "institution": "Kabul Medical University"}],
        "clinical_experience": {"years": 3},
        "languages": [{"name": "Dari"}],
        "preferences": {"locations": ["Kabul"]},
    }
    report = match_job_against_profile(job, profile).to_dict()
    docs = generate_tailored_documents(job, profile, report)
    assert "professional license" not in docs["cover_letter"].lower()
    assert any("License" in warning or "license" in warning for warning in docs["review_warnings"])


def test_application_state_transitions_require_confirmation(tmp_path, monkeypatch):
    monkeypatch.setattr(tracker, "DB_PATH", tmp_path / "apps.db")
    job = Job("job-state", "Medical Officer", "NGO", "Kabul", "https://x/job", "https://x/apply", "test")
    tracker.log_discovered(job)
    ok, msg = tracker.transition_application_state("job-state", "opened")
    assert ok, msg
    ok, msg = tracker.transition_application_state("job-state", "submitted", explicit_confirmation=False)
    assert not ok
    assert "confirmation" in msg.lower()
    ok, msg = tracker.transition_application_state("job-state", "submitted", explicit_confirmation=True)
    assert ok, msg
    assert tracker.get_job_by_id("job-state")["status"] == "submitted"


def test_low_priority_job_can_still_be_prepared_and_opened_for_review(tmp_path, monkeypatch):
    import utils.tracker as tracker
    from utils.discovery import Job

    monkeypatch.setattr(tracker, "DB_PATH", tmp_path / "applications.db")
    job = Job(
        id="low_priority_review_job",
        title="Clinic Manager",
        company="Health NGO",
        location="Kabul",
        url="https://example.org/job",
        apply_url="https://example.org/apply",
        platform="test",
        description="",
    )
    tracker.log_discovered(job)
    tracker.log_medical_match(job.id, {"priority": "Low priority", "explanation": "Not enough years", "facts": {}})
    assert tracker.get_job_by_id(job.id)["status"] == "not_eligible"
    tracker.update_tailored_resume(job.id, {"tailored_cv_text": "draft", "cover_letter": "letter"})
    assert tracker.get_job_by_id(job.id)["status"] == "prepared"
    ok, message = tracker.transition_application_state(job.id, "opened")
    assert ok, message
    ok, message = tracker.transition_application_state(job.id, "submitted", explicit_confirmation=False)
    assert not ok
    assert "Explicit user confirmation" in message


def test_tailored_cv_is_application_ready_and_keeps_warnings_separate():
    job = {
        "id": "j-review-ready",
        "title": "Medical Officer",
        "company": "Health NGO",
        "location": "Kabul",
        "url": "https://example.org/job",
        "apply_url": "https://example.org/apply",
        "description": "Medical Officer. MD and valid medical license required. Closing date: 30 September 2026.",
        "metadata": {},
    }
    profile = {
        "personal": {"first_name": "A", "last_name": "Doctor", "email": "a@example.org", "location": "Kabul"},
        "medical_education": [{"degree": "MD", "institution": "Kabul Medical University"}],
        "clinical_experience": {"years": 3},
        "work_history": [{"title": "Medical Officer", "organization": "Clinic", "start": "2024-01", "end": "2024-12", "bullets": ["Provided clinical care."]}],
        "languages": [{"name": "Dari"}],
        "preferences": {"locations": ["Kabul"]},
    }
    report = match_job_against_profile(job, profile).to_dict()
    docs = generate_tailored_documents(job, profile, report)
    cv_lower = docs["tailored_cv_text"].lower()
    cover_lower = docs["cover_letter"].lower()
    assert "tailored cv draft" not in cv_lower
    assert "review notes" not in cv_lower
    assert "needs verification" not in cv_lower
    assert "i have not included" not in cover_lower
    assert any("license" in warning.lower() for warning in docs["review_warnings"])


def test_tailored_cv_core_competencies_dedupe_nested_acronyms():
    job = {
        "id": "j-dedupe",
        "title": "Quality Officer",
        "company": "Health NGO",
        "location": "Kabul",
        "url": "https://example.org/job",
        "apply_url": "https://example.org/apply",
        "description": "MD required. BPHS EPHS IMAM HMIS required. Apply at https://example.org/apply. Deadline: 2026-10-10.",
        "metadata": {},
    }
    profile = {
        "personal": {"first_name": "A", "last_name": "Doctor"},
        "medical_education": [{"degree": "MD"}],
        "skills": {"public_health": ["BPHS/EPHS & IMAM/CMAM Program Knowledge", "BPHS", "EPHS", "IMAM", "HMIS"]},
        "languages": [{"name": "Dari"}],
    }
    report = match_job_against_profile(job, profile).to_dict()
    docs = generate_tailored_documents(job, profile, report)
    core_line = next(line for line in docs["tailored_cv_text"].splitlines() if line.startswith("- "))
    assert "BPHS/EPHS & IMAM/CMAM Program Knowledge, BPHS" not in core_line
    assert "BPHS/EPHS & IMAM/CMAM Program Knowledge, EPHS" not in core_line


def test_email_application_package_includes_complete_draft_and_checklists():
    job = {
        "id": "email-ready",
        "title": "TSFP Project Supervisor",
        "company": "HealthNet TPO",
        "location": "Kunar",
        "url": "https://www.acbar.org/en/jobs/details/145882/tsfp-project-supervisor",
        "apply_url": "mailto:recruitment.kabul@hntpo.org",
        "description": "Send CV and application letter to recruitment.kabul@hntpo.org. Mention the position in the subject line. Deadline: 2026-10-10.",
        "metadata": {
            "source_url": "https://www.acbar.org/en/jobs/details/145882/tsfp-project-supervisor",
            "required_documents": ["Tailored CV", "Tailored cover letter"],
            "special_instructions": ["Mention the position in the email subject line"],
        },
    }
    profile = {
        "personal": {"first_name": "Allah Yar", "last_name": "Frotan", "email": "alfrotan@example.org", "phone": "+93 700 000 000", "location": "Kabul"},
        "medical_education": [{"degree": "MD", "institution": "Kabul Medical Science University"}],
        "license_registration": {"verified": True, "status": "Verified — holds valid medical professional registration/license", "number": ""},
        "languages": [{"name": "Dari"}, {"name": "English"}],
    }
    match_report = {
        "facts": {
            "source_url": job["url"],
            "application_email": "recruitment.kabul@hntpo.org",
            "application_subject": "TSFP Project Supervisor",
            "application_subject_required": True,
            "closing_date": "2026-10-10",
        },
        "requirement_matches": [],
    }
    docs = {"suggested_subject": "TSFP Project Supervisor"}
    package = generate_application_package(
        job,
        profile,
        match_report,
        docs,
        generated_paths={"tailored_cv": "cv.txt", "cover_letter": "cover.txt"},
    )
    assert package["package_status"] == "READY_TO_SEND"
    assert package["route_type"] == "email"
    assert package["email_draft"]["to"] == "recruitment.kabul@hntpo.org"
    assert package["email_draft"]["subject"] == "TSFP Project Supervisor"
    assert "Dear Hiring Committee" in package["email_draft"]["body"]
    assert "cv.txt" in package["email_draft"]["attachments"]
    assert "cover.txt" in package["email_draft"]["attachments"]
    assert package["source_url"] == job["url"]
    assert package["deadline"] == "2026-10-10"
    assert "license number:" not in package["text"].lower()


def test_cover_letter_does_not_claim_license_requirement_when_vacancy_does_not_require_it():
    job = {
        "id": "health-community-role",
        "title": "Social Mobilizer",
        "company": "Health NGO",
        "location": "Kunar",
        "url": "https://example.org/job",
        "apply_url": "https://forms.gle/example",
        "description": "Relevant education in public health. Mobilize communities to utilize health and nutrition services. Deadline: 2026-10-10.",
        "metadata": {},
    }
    profile = {
        "personal": {"first_name": "Allah Yar", "last_name": "Frotan", "location": "Kabul, Afghanistan"},
        "medical_education": [{"degree": "Doctor of Medicine (MD)", "institution": "Kabul Medical Science University"}],
        "license_registration": {"verified": True, "status": "Verified — holds valid medical professional registration/license", "number": ""},
        "work_history": [{"title": "Health and Nutrition Supervisor", "organization": "ACF", "start": "2022-02", "end": "2022-12", "bullets": ["Supervised mobile health and nutrition teams delivering community-based BPHS/EPHS-related services."]}],
        "skills": {"public_health": ["Health & Nutrition Program Supervision", "BPHS", "EPHS", "nutrition"]},
        "languages": [{"name": "Dari"}, {"name": "English"}],
    }
    report = match_job_against_profile(job, profile).to_dict()
    docs = generate_tailored_documents(job, profile, report)
    assert "registration/license requirement stated for the role" not in docs["cover_letter"]


def test_application_package_surfaces_unverified_essential_requirements():
    job = {
        "id": "surgeon-needs-verification",
        "title": "Surgeon",
        "company": "AKHS-A",
        "location": "Kabul",
        "url": "https://akhs.odoo.com/jobs/detail/surgeon-125",
        "apply_url": "https://akhs.odoo.com/jobs/detail/surgeon-125",
        "description": "Specialist Surgeon qualification and specialist certification in general surgery required. Deadline: 2026-10-10.",
        "metadata": {},
    }
    profile = {
        "personal": {"first_name": "Allah Yar", "last_name": "Frotan", "location": "Kabul"},
        "medical_education": [{"degree": "MD"}],
        "work_history": [{"title": "Medical Doctor", "organization": "Clinic", "start": "2022-01", "end": "2025-07", "bullets": ["Provided clinical care."]}],
        "languages": [{"name": "Dari"}, {"name": "Pashto"}],
    }
    report = match_job_against_profile(job, profile, today=date(2026, 9, 30)).to_dict()
    docs = generate_tailored_documents(job, profile, report)
    package = generate_application_package(job, profile, report, docs)
    assert package["package_status"] == "NEEDS_USER_INPUT"
    assert any("Specialist" in item or "specialist" in item for item in package["missing_items"])
    assert any("Specialist" in item or "specialist" in item for item in package["review_warnings"])
    assert "Verification warnings before submission" in package["text"]
    assert "specialist" in package["text"].lower()
    assert "Submit only after explicit user confirmation" in package["text"]



def test_placeholder_contact_values_are_flagged_in_packages():
    job = {
        "id": "placeholder-contact",
        "title": "Medical Officer",
        "company": "FMIC",
        "location": "Kabul",
        "url": "https://example.org/job",
        "apply_url": "https://example.org/apply",
        "description": "Medical Officer required. Deadline: 2026-10-05.",
        "metadata": {},
    }
    profile = {
        "personal": {"first_name": "Allah Yar", "last_name": "Frotan", "email": "doctor@example.org", "phone": "+93 000 000 000", "location": "Kabul"},
        "medical_education": [{"degree": "MD"}],
        "languages": [{"name": "Dari", "level": "Native"}, {"name": "English", "level": "Fluent"}, {"name": "Pashto", "level": "Intermediate"}],
    }
    report = {"facts": {"source_url": job["url"], "application_url": job["apply_url"], "closing_date": "2026-10-05"}, "requirement_matches": []}
    docs = generate_tailored_documents(job, profile, report)
    package = generate_application_package(job, profile, report, docs)
    assert "doctor@example.org" not in docs["tailored_cv_text"]
    assert "+93 000 000 000" not in docs["tailored_cv_text"]
    assert "CONFIRM BEFORE SUBMISSION" in docs["tailored_cv_text"]
    assert "Pashto — Intermediate" in docs["tailored_cv_text"]
    assert package["package_status"] == "NEEDS_USER_INPUT"
    assert "Confirmed personal email address" in package["missing_items"]
    assert "Confirmed phone number" in package["missing_items"]


def test_online_application_package_includes_form_url_fields_and_no_submission():
    job = {
        "id": "form-ready",
        "title": "Medical Officer",
        "company": "FMIC",
        "location": "Kabul",
        "url": "https://www.acbar.org/en/jobs/details/145774/medical-officer",
        "apply_url": "https://docs.google.com/forms/d/example/edit",
        "description": "Apply through the online form. Validated copies of academic certificates/diplomas will be requested only if selected. Deadline: 2026-10-05.",
        "metadata": {
            "source_url": "https://www.acbar.org/en/jobs/details/145774/medical-officer",
            "required_documents": ["Resume/CV with specific qualification and experience dates", "Academic certificates/diplomas only if selected/requested"],
            "form_fields": ["Full name", "Phone", "Email", "Work history with dates", "CV upload"],
        },
    }
    profile = {
        "personal": {"first_name": "Allah Yar", "last_name": "Frotan", "email": "alfrotan@example.org", "phone": "+93 700 000 000", "location": "Kabul"},
        "medical_education": [{"degree": "MD"}],
        "license_registration": {"verified": True, "status": "Verified — holds valid medical professional registration/license", "number": ""},
    }
    match_report = {"facts": {"source_url": job["url"], "application_url": job["apply_url"], "closing_date": "2026-10-05"}, "requirement_matches": []}
    package = generate_application_package(job, profile, match_report, {"suggested_subject": "Medical Officer"})
    assert package["package_status"] == "READY_TO_SUBMIT"
    assert package["route_type"] == "online_form"
    assert package["online_application"]["url"] == "https://docs.google.com/forms/d/example/viewform"
    assert package["application_route"] == "https://docs.google.com/forms/d/example/viewform"
    assert any("/edit URL" in item for item in package["special_instructions"])
    assert "CV upload" in package["online_application"]["form_fields_checklist"]
    assert package["email_draft"] is None
    assert package["no_submission_performed"] is True
    assert any("CAPTCHA" in item for item in package["user_required_actions"])
    assert "license number:" not in package["text"].lower()


def test_email_application_package_blocks_send_when_required_external_form_missing():
    job = {
        "id": "email-blocked",
        "title": "Quality Officer",
        "company": "RI",
        "location": "Nimruz",
        "url": "https://example.org/job",
        "apply_url": "mailto:vacancies@example.org",
        "description": "Send RI application form with CV and three references. Subject line: REF-1.",
        "metadata": {
            "required_documents": ["RI application form", "CV", "Three references"],
            "form_fields": ["Father’s name", "Date of birth", "Tazkera number", "Three references"],
        },
    }
    profile = {"personal": {"first_name": "Allah Yar", "last_name": "Frotan"}}
    report = {"facts": {"application_email": "vacancies@example.org", "application_subject": "REF-1", "application_subject_required": True}}
    package = generate_application_package(job, profile, report, {"suggested_subject": "REF-1"})
    assert package["package_status"] == "NEEDS_USER_INPUT"
    assert "Completed organization application form" in package["missing_items"]
    assert "Approved professional reference contacts" in package["missing_items"]
    assert "Father’s name" in package["form_fields_checklist"]
    assert "Three references" in package["text"]
    assert any("Approved professional reference contacts" in action for action in package["user_required_actions"])


def _technical_supervisor_profile():
    return {
        "personal": {"first_name": "Allah Yar", "last_name": "Frotan", "email": "alfrotan@example.org", "phone": "+93 700 000 000", "location": "Kabul, Afghanistan"},
        "medical_education": [{"degree": "Doctor of Medicine (MD), Curative Medicine", "institution": "Kabul Medical Science University", "start": "2013", "end": "2020"}],
        "license_registration": {"verified": True, "status": "Verified — holds valid medical professional registration/license", "number": ""},
        "work_history": [
            {
                "title": "TFU Medical Doctor & Safeguarding Focal Point",
                "organization": "Action Against Hunger (ACF-International)",
                "location": "Daikundi",
                "start": "2023-05",
                "end": "2025-07",
                "bullets": [
                    "Led day-to-day Therapeutic Feeding Unit (TFU) operations, coordinating closely with nursing staff, nutrition teams, and auxiliary personnel to ensure continuity of inpatient care.",
                    "Maintained accurate HMIS records and served as Safeguarding Focal Point supporting PSEA and child protection.",
                    "Coordinated medical supply forecasting and stock monitoring to safeguard uninterrupted availability of essential medical items.",
                ],
            },
            {
                "title": "Health and Nutrition Supervisor",
                "organization": "Action Against Hunger (ACF-International)",
                "location": "Daikundi",
                "start": "2022-02",
                "end": "2022-12",
                "bullets": [
                    "Supervised mobile health and nutrition teams delivering community-based BPHS/EPHS-related services.",
                    "Monitored field-level program activities, identified service gaps, and supported corrective actions to improve program quality and coverage.",
                ],
            },
            {
                "title": "Medical Doctor / COVID-19 Rapid Response Team Leader",
                "organization": "Daikundi Provincial Public Health Directorate",
                "location": "Daikundi",
                "start": "2020-10",
                "end": "2020-12",
                "bullets": ["Coordinated with Ministry of Public Health representatives, provincial health authorities, and partner organizations."],
            },
        ],
        "skills": {"public_health": ["BPHS/EPHS & IMAM/CMAM Program Knowledge", "HMIS / DHIS2 Reporting & Data Quality"], "management": ["Health & Nutrition Program Supervision"]},
        "languages": [{"name": "Dari", "level": "Native"}, {"name": "English", "level": "Fluent"}, {"name": "Pashto", "level": "Intermediate"}],
        "certificates": ["Infection Prevention & Control (IPC) — ACF (2023)", "Project Management — Coventry University, UK (2023)"],
        "preferences": {"locations": ["Afghanistan", "Kabul", "Daikundi"], "field_deployment": True, "willing_to_relocate": True},
    }


def test_tailored_documents_are_vacancy_specific_for_technical_supervisor():
    job = {
        "id": "145872",
        "title": "Technical Supervisor",
        "company": "Bakhter Development Network - BDN",
        "location": "Takhar",
        "url": "https://www.acbar.org/en/jobs/details/145872/technical-supervisor",
        "apply_url": "https://forms.gle/U3dJcBHhJamUZ5yE9",
        "description": """
        Medical degree (MD) and a valid medical license to practice in Afghanistan.
        Minimum of 3 years of experience in healthcare project management. Strong knowledge of primary healthcare principles and practices.
        Supervising primary healthcare services, strengthening patient referral systems, coordinating medicines and medical supplies,
        providing technical guidance, supportive supervision and mentorship, maintaining records and submitting reports, English fluency.
        Closing date: 08 October 2026. Application Form: https://forms.gle/U3dJcBHhJamUZ5yE9
        """,
        "metadata": {"source_url": "https://www.acbar.org/en/jobs/details/145872/technical-supervisor"},
    }
    profile = _technical_supervisor_profile()
    report = match_job_against_profile(job, profile, today=date(2026, 9, 30)).to_dict()
    docs = generate_tailored_documents(job, profile, report)
    cv = docs["tailored_cv_text"]
    letter = docs["cover_letter"]
    assert "TARGET ROLE: Technical Supervisor" in cv
    assert "VACANCY-FIT HIGHLIGHTS" in cv
    assert "medical supply forecasting" in cv
    assert "Supervised mobile health and nutrition teams" in cv
    assert "I understand that the vacancy emphasizes" in letter
    assert "medical supply forecasting" in letter
    assert "license number" not in (cv + letter).lower()
    assert "registration number" not in (cv + letter).lower()


def test_prepare_application_bundle_writes_complete_review_files(tmp_path):
    job = {
        "id": "145872",
        "title": "Technical Supervisor",
        "company": "Bakhter Development Network - BDN",
        "location": "Takhar",
        "url": "https://www.acbar.org/en/jobs/details/145872/technical-supervisor",
        "apply_url": "https://forms.gle/U3dJcBHhJamUZ5yE9",
        "description": "Medical degree MD, valid medical license, 3 years healthcare project management, PHC, supportive supervision, reports, English. Closing date: 08 October 2026. Application Form: https://forms.gle/U3dJcBHhJamUZ5yE9",
        "metadata": {"required_documents": ["Updated CV", "Formal cover letter"], "form_fields": ["Full name", "Email", "CV upload", "Cover letter upload"]},
    }
    profile = _technical_supervisor_profile()
    report = match_job_against_profile(job, profile, today=date(2026, 9, 30)).to_dict()
    docs = prepare_application_bundle(job, profile, report, out_dir=tmp_path)
    paths = docs["generated_paths"]
    assert Path(paths["tailored_cv"]["txt"]).exists()
    assert Path(paths["tailored_cv"]["docx"]).exists()
    assert Path(paths["tailored_cv"]["pdf"]).exists()
    assert Path(paths["cover_letter"]["txt"]).exists()
    assert Path(paths["cover_letter"]["docx"]).exists()
    assert Path(paths["cover_letter"]["pdf"]).exists()
    assert Path(paths["application_package_txt"]).exists()
    package_text = Path(paths["application_package_txt"]).read_text(encoding="utf-8")
    assert "Package status: READY_TO_SUBMIT" in package_text
    assert "Application URL/form: https://forms.gle/U3dJcBHhJamUZ5yE9" in package_text
    assert "Updated CV" in package_text
    assert "Formal cover letter" in package_text
    assert "Submit only after explicit user confirmation" in package_text
    assert docs["no_submission_performed"] is True


def test_global_document_design_system_renders_cv_and_cover_letter_for_every_bundle(tmp_path):
    import fitz
    from docx import Document
    from utils.document_design import DESIGN_SYSTEM_VERSION

    profile = _technical_supervisor_profile()
    jobs = [
        {
            "id": "watch_aeb9c295fd34296adf",
            "title": "Medical Officer",
            "company": "French Medical Institute for Mothers and Children (FMIC)",
            "location": "Kabul",
            "url": "https://www.acbar.org/en/jobs/details/145774/medical-officer",
            "apply_url": "https://docs.google.com/forms/d/example/viewform",
            "description": "Medical Officer, MD, valid medical registration, clinical care, inpatient/outpatient service, quality and patient safety. Vacancy reference FMIC/HR/571. Closing date: 05 October 2026.",
            "metadata": {"reference_number": "FMIC/HR/571", "required_documents": ["Updated CV", "Cover letter"]},
        },
        {
            "id": "145872",
            "title": "Technical Supervisor",
            "company": "Bakhter Development Network - BDN",
            "location": "Takhar",
            "url": "https://www.acbar.org/en/jobs/details/145872/technical-supervisor",
            "apply_url": "https://forms.gle/U3dJcBHhJamUZ5yE9",
            "description": "Medical degree MD, valid medical license, 3 years healthcare project management, PHC, supportive supervision, reports, English. Closing date: 08 October 2026.",
            "metadata": {"required_documents": ["Updated CV", "Formal cover letter"]},
        },
    ]
    generated = []
    for job in jobs:
        report = match_job_against_profile(job, profile, today=date(2026, 9, 30)).to_dict()
        docs = prepare_application_bundle(job, profile, report, out_dir=tmp_path)
        paths = docs["generated_paths"]
        generated.append((job, docs, paths))

        txt = Path(paths["tailored_cv"]["txt"]).read_text(encoding="utf-8")
        assert "TARGET ROLE:" in txt
        assert DESIGN_SYSTEM_VERSION not in txt  # ATS/plain-text output is not polluted by layout markers.

        cv_pdf = fitz.open(paths["tailored_cv"]["pdf"])
        cover_pdf = fitz.open(paths["cover_letter"]["pdf"])
        assert cv_pdf.page_count >= 2
        assert cover_pdf.page_count == 1
        cv_text = "\n".join(page.get_text() for page in cv_pdf)
        cover_text = "\n".join(page.get_text() for page in cover_pdf)
        assert "APPLICATION FOCUS" in cv_text
        assert "PROFESSIONAL EXPERIENCE" in cv_text
        assert job["title"] in cv_text
        assert job["company"].split(" (")[0] in cv_text
        assert "APPLICATION LETTER" in cover_text
        assert job["title"] in cover_text
        assert job["company"].split(" - ")[0].split(" (")[0] in cover_text

        cv_docx = Document(paths["tailored_cv"]["docx"])
        cover_docx = Document(paths["cover_letter"]["docx"])
        cv_style_names = {style.name for style in cv_docx.styles}
        assert {"JF Name", "JF Section", "JF Body"}.issubset(cv_style_names)
        cv_docx_text = "\n".join(
            [paragraph.text for paragraph in cv_docx.paragraphs]
            + [paragraph.text for table in cv_docx.tables for row in table.rows for cell in row.cells for paragraph in cell.paragraphs]
        )
        cover_docx_text = "\n".join(paragraph.text for paragraph in cover_docx.paragraphs)
        assert "PROFESSIONAL EXPERIENCE" in cv_docx_text
        assert "APPLICATION LETTER" in cover_docx_text

    fmic_cv = Path(generated[0][2]["tailored_cv"]["txt"]).read_text(encoding="utf-8")
    bdn_cv = Path(generated[1][2]["tailored_cv"]["txt"]).read_text(encoding="utf-8")
    assert "TARGET ORGANIZATION: French Medical Institute for Mothers and Children" in fmic_cv
    assert "TARGET ORGANIZATION: Bakhter Development Network - BDN" in bdn_cv
    assert fmic_cv != bdn_cv
