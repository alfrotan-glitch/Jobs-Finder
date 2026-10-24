"""Acceptance contract: professional, job-targeted, factual CV design.

The actual Head of TFU applicant's canonical professional evidence is tested
without fabricating any outcomes, managerial promotions or numeric achievements.
The generated TXT, Word and PDF must agree on the role being APPLIED FOR while
retaining the applicant's actual, verified professional title.
"""

from datetime import date
from pathlib import Path

import pdfplumber
from docx import Document

from utils.document_design import render_professional_document_artifacts
from utils.documents import generate_master_cv, generate_tailored_documents
from utils.medical_matcher import match_job_against_profile
from utils.profile import load_canonical_profile


def _head_tfu_job():
    return {
        "id": "editorial-head-tfu",
        "title": "Head of TFU",
        "company": "Illustrative Health NGO",
        "location": "Daikundi, Afghanistan",
        "url": "https://example.org/vacancies/head-tfu",
        "apply_url": "https://example.org/apply/head-tfu",
        "description": (
            "Head of Therapeutic Feeding Unit (TFU). Supervise clinical "
            "treatment of children with severe acute malnutrition (SAM), "
            "IMAM/CMAM services, stock and programme reporting, team coaching "
            "and nutrition quality assurance. Medical Doctor accepted."
        ),
        "metadata": {
            "source_name": "ACBAR", "source_url": "https://example.org/jobs",
            "closing_date": "2026-12-31",
        },
    }


def test_head_of_tfu_targeted_cv_has_clear_position_label_without_fake_appointment(tmp_path):
    profile = load_canonical_profile(required=True)
    job = _head_tfu_job()
    report = match_job_against_profile(job, profile, today=date(2026, 10, 20)).to_dict()
    tailored = generate_tailored_documents(job, profile, report)
    model = tailored["tailored_cv_model"]
    assert model["application_focus"] == "Head of TFU"
    assert model["headline"] != "Head of TFU"
    assert "therapeutic feeding unit (TFU) clinical care" in model["profile"]
    assert "health and nutrition team supervision" in model["profile"]
    assert "TFU Medical Doctor & Safeguarding Focal Point" not in model["profile"]
    assert "Head of TFU" not in model["profile"]
    assert "Relevant experience includes" in model["profile"]
    assert tailored["tailored_cv_text"].count("APPLICATION FOR: Head of TFU") == 1
    assert "serving as Head of TFU" not in tailored["tailored_cv_text"]

    paths = render_professional_document_artifacts(
        tailored["tailored_cv_text"], tmp_path / "head_tfu_cv", document_type="cv",
        metadata={"job": job}, canonical_cv_model=model,
    )
    txt = Path(paths["txt"]).read_text(encoding="utf-8")
    doc = Document(paths["docx"])
    word = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    with pdfplumber.open(paths["pdf"]) as pdf:
        pdf_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        assert 1 <= len(pdf.pages) <= 3
        assert any(
            any("APPLICATION FOR" in line.get("text", "") for line in page.extract_text_lines())
            for page in pdf.pages
        )
    for text in (txt, word, pdf_text):
        assert "Head of TFU" in text
        assert "TFU Medical Doctor & Safeguarding Focal Point" in text
        assert "Health & Nutrition Supervisor" in text
        assert "ACF-International" in text
        assert "Medical Doctor" in text
    assert "APPLICATION FOR  Head of TFU" in word
    assert not doc.tables
    assert [p.text for p in doc.paragraphs if p.style.name == "CV Application Focus"] == [
        "APPLICATION FOR  Head of TFU"
    ]


def test_master_cv_remains_neutral_and_unsupported_outcomes_stay_out():
    profile = load_canonical_profile(required=True)
    master = generate_master_cv(profile)
    model = master["master_cv_model"]
    assert "application_focus" not in model
    assert "APPLICATION FOR" not in master["master_cv_text"]
    text = master["master_cv_text"]
    assert "Head of TFU" not in text
    for unsupported in ("reduced mortality by", "increased admissions by 50%", "managed 100 staff"):
        assert unsupported not in text.lower()
