"""Real-output regressions: verified data -> bundle -> searchable print layout.

Geometry checks supplement (not replace) the page-by-page visual audit recorded
in docs/cv_rendering_audit.md. No golden PDF or applicant artifact is committed.
"""
from __future__ import annotations

import copy
import re
from datetime import date
from itertools import pairwise
from pathlib import Path

import pdfplumber
import pytest
from docx import Document
from docx.oxml.ns import qn

from utils import document_design as design
from utils.documents import (
    _experience_dates,
    _reverse_chronological_entries,
    generate_master_cv,
    prepare_application_bundle,
    write_master_cv,
)
from utils.medical_matcher import match_job_against_profile
from utils.profile import canonical_profile_path, load_canonical_profile

SECTIONS = [
    "PROFESSIONAL SUMMARY", "CORE PROFESSIONAL COMPETENCIES",
    "PROFESSIONAL EXPERIENCE", "EDUCATION", "PROFESSIONAL REGISTRATION",
    "MEDICAL EXIT EXAMINATION", "PROFESSIONAL TRAINING & CERTIFICATIONS",
    "LANGUAGES", "REFERENCES",
]
VACANCIES = [
    ("Medical Officer", "Clinical care, HMIS reporting and medical registration"),
    ("Health and Nutrition Supervisor", "IMAM/CMAM, SAM, IYCF and supervision"),
    ("Provincial Public Health Coordinator", "MoPH coordination, BPHS/EPHS and emergency response"),
    ("Safeguarding Officer", "Safeguarding, PSEA and child protection"),
]


def vacancy(index):
    title, focus = VACANCIES[index]
    return {
        "id": f"cv-render-{index}", "title": title, "company": "Example Health NGO",
        "location": "Afghanistan", "url": f"https://example.org/vacancy/{index}",
        "apply_url": "recruitment@example.org",
        "description": f"Medical Doctor accepted. {focus} required. Apply to recruitment@example.org by 2026-12-31.",
        "metadata": {"closing_date": "2026-12-31"},
    }


def flat(text):
    return " ".join(text.split())


def assert_geometry(pdf):
    """Detect crop/overflow, touching text lines, footer collisions and bad pages."""
    for number, page in enumerate(pdf.pages, 1):
        assert page.width == pytest.approx(595.276, abs=.1)
        assert page.height == pytest.approx(841.890, abs=.1)
        assert all(49 <= c["x0"] < c["x1"] <= page.width - 49 for c in page.chars if c["text"].strip())
        assert all(43 <= c["top"] < c["bottom"] <= page.height - 16 for c in page.chars)
        footer = page.crop((0, page.height - 37, page.width, page.height)).extract_text()
        assert re.findall(r"Page \d+", footer) == [f"Page {number}"]
        body = page.crop((0, 0, page.width, page.height - 45))
        lines = body.extract_text_lines()
        assert lines
        # All CV content is one flowing column, so successive lines cannot
        # overlap vertically. Glyph boxes include ascenders/descenders.
        for previous, following in pairwise(lines):
            assert previous["bottom"] <= following["top"] + .5, (previous, following)
        assert lines[-1]["bottom"] < page.height - 47
        # No single section heading left hanging at the bottom of a page.
        assert lines[-1]["text"] not in SECTIONS
        # Rules must really exist in the PDF, not just in ParagraphStyle kwargs.
        for line in lines:
            if line["text"] in SECTIONS:
                assert any(line["bottom"] <= rule["top"] <= line["bottom"] + 8 for rule in page.lines)
                assert "Bold" in line["chars"][0]["fontname"]
    if len(pdf.pages) > 1:
        last = pdf.pages[-1]
        chars = [c for c in last.chars if c["bottom"] < last.height - 45]
        assert (max(c["bottom"] for c in chars) - min(c["top"] for c in chars)) / (last.height - 90) >= .4


def assert_model_preserved(model, text):
    text = flat(text)
    values = [model["name"], model["headline"], model["profile"], *model["contact_lines"]]
    for key in ("education", "registration", "exit_exam", "certifications", "references"):
        values.extend(model[key])
    for group in model["expertise"]:
        values.extend([group["group"], *group["items"]])
    for name, level in model["languages"]:
        values.extend([name, level])
    for role in model["experience"]:
        values.extend([role["role"], role["org"], role["loc"], role["dates"], *role["bullets"]])
    for value in values:
        assert flat(value) in text, value
    positions = [text.index(role["role"]) for role in model["experience"]]
    assert positions == sorted(positions)
    positions = [text.index(heading) for heading in SECTIONS]
    assert positions == sorted(positions)


@pytest.mark.parametrize("kind", ["master", 0, 1, 2, 3])
def test_real_cv_complete_readable_and_immutable_in_every_format(tmp_path, kind):
    profile = load_canonical_profile(required=True)
    before = copy.deepcopy(profile)
    source_bytes = canonical_profile_path().read_bytes()
    master = write_master_cv(profile, out_dir=tmp_path / "master")
    master_bytes = {key: Path(path).read_bytes() for key, path in master["generated_paths"].items()}
    if kind == "master":
        model, paths = master["master_cv_model"], master["generated_paths"]
    else:
        job = vacancy(kind)
        report = match_job_against_profile(job, profile, today=date(2026, 10, 9)).to_dict()
        bundle = prepare_application_bundle(job, profile, report, out_dir=tmp_path / "tailored")
        model, paths = bundle["tailored_cv_model"], bundle["generated_paths"]["tailored_cv"]
    assert [r["role"] for r in model["experience"]] == [
        "TFU Medical Doctor & Safeguarding Focal Point", "Health & Nutrition Supervisor",
        "Public Relations & Communications Advisor", "Medical Doctor / COVID-19 Rapid Response Team Leader",
        "Administrative and Finance Officer",
    ]
    assert_model_preserved(model, Path(paths["txt"]).read_text(encoding="utf-8"))
    doc = Document(paths["docx"])
    assert_model_preserved(model, "\n".join(p.text for p in doc.paragraphs))
    with pdfplumber.open(paths["pdf"]) as pdf:
        assert len(pdf.pages) == 2  # this real evidence set, not an exporter cap
        assert_model_preserved(model, "\n".join(p.extract_text() for p in pdf.pages))
        assert_geometry(pdf)
        for role in model["experience"]:
            # Entire short role header and first responsibility stay together.
            assert any(all(flat(v) in flat(p.extract_text()) for v in [role["role"], role["org"], role["loc"], role["dates"], role["bullets"][0]]) for p in pdf.pages)
    assert profile == before
    assert canonical_profile_path().read_bytes() == source_bytes
    assert generate_master_cv(profile)["master_cv_text"] == master["master_cv_text"]
    for key, path in master["generated_paths"].items():
        assert Path(path).read_bytes() == master_bytes[key]


def test_word_shared_typography_semantic_hierarchy_and_pagination(tmp_path):
    master = write_master_cv(load_canonical_profile(required=True), out_dir=tmp_path)
    doc = Document(master["generated_paths"]["docx"])
    section = doc.sections[0]
    assert section.page_width.mm == pytest.approx(210, abs=.1)
    assert section.page_height.mm == pytest.approx(297, abs=.1)
    assert section.left_margin.mm == pytest.approx(design.CV_MARGIN_MM, abs=.1)
    assert section.right_margin == section.left_margin
    for name in ("CV Body", "CV Bullet", "CV Competency Group"):
        st = doc.styles[name]
        assert st.font.size.pt == pytest.approx(design.CV_PDF_LAYOUTS["comfortable"]["body"], abs=.1)
        assert st.font.name == design._cv_word_font()
        assert st.paragraph_format.keep_together
        assert st.paragraph_format.widow_control
    assert doc.styles["CV Section"].base_style.name == "Heading 1"
    assert doc.styles["CV Role"].base_style.name == "Heading 2"
    for name in ("CV Section", "CV Role", "CV Meta"):
        assert doc.styles[name].paragraph_format.keep_with_next
    assert [p.text for p in doc.paragraphs if p.style.name == "CV Section"] == SECTIONS
    assert not doc.tables  # no text boxes, floating panels or layout tables
    assert not doc.element.xpath(".//w:drawing | .//w:txbxContent | .//w:br[@w:type='page']")
    assert all(p._p.xpath("./w:pPr/w:pBdr/w:bottom") for p in doc.paragraphs if p.style.name == "CV Section")
    footer = section.footer.paragraphs[0]
    tab = next(iter(footer.paragraph_format.tab_stops))
    assert tab.position == section.page_width - section.left_margin - section.right_margin
    assert footer._p.xpath(".//w:instrText")[0].text == "PAGE"
    assert doc.settings.element.find(qn("w:updateFields")).get(qn("w:val")) == "true"


def test_chronology_is_stable_and_does_not_turn_unknown_dates_into_current_jobs():
    entries = [
        {"title": "Undated Z"}, {"title": "Old", "end": "2020"},
        {"title": "Current", "start": "2021", "end": "Present"},
        {"title": "Undated A", "end": "unconfirmed"},
        {"title": "Start only", "start": "2022-02"},
        {"title": "Latest completed", "start": "2023-05", "end": "2025-07"},
        {"title": "Same date", "start": "2023-05", "end": "2025-07"},
    ]
    before = copy.deepcopy(entries)
    ordered = _reverse_chronological_entries(entries)
    assert [r["title"] for r in ordered] == ["Current", "Latest completed", "Same date", "Start only", "Old", "Undated Z", "Undated A"]
    assert entries == before
    assert _experience_dates(entries[0]) == ""
    assert _experience_dates(entries[1]) == "2020"
    assert _experience_dates(entries[4]) == "Feb 2022"
    assert _experience_dates(entries[5]) == "May 2023 – Jul 2025"


def test_long_role_headers_and_multi_page_history_do_not_clip_or_orphan(tmp_path):
    # Presentation fixture only: never inserted into the applicant record.
    model = {
        "name": "Rendering Test", "headline": "Layout fixture",
        "experience": [
            {"role": f"Role {i:02d} — " + "Long clinical programme title " * 4,
             "org": f"Employer {i:02d} — " + "Long organization name " * 3,
             "loc": "A deliberately long location description for wrapping checks",
             "dates": "2020 – 2021",
             "bullets": [f"Unique responsibility {i:02d}: " + "Verified fixture content for pagination testing. " * 6]}
            for i in range(12)
        ],
    }
    path = tmp_path / "long.pdf"
    design.render_cv_pdf(model, path)
    with pdfplumber.open(path) as pdf:
        assert len(pdf.pages) >= 3
        assert_geometry(pdf)
        text = flat(" ".join(p.extract_text() for p in pdf.pages))
        for role in model["experience"]:
            for value in [role["role"], role["org"], *role["bullets"]]:
                assert flat(value) in text
            assert any(flat(role["role"]) in flat(p.extract_text()) and flat(role["bullets"][0][:50]) in flat(p.extract_text()) for p in pdf.pages)


def test_whitespace_compaction_never_shrinks_text_and_restores_if_ineffective(tmp_path, monkeypatch):
    for key in ("name", "title", "contact", "section", "role", "meta", "body", "bullets"):
        assert design.CV_PDF_LAYOUTS["compact"][key] == design.CV_PDF_LAYOUTS["comfortable"][key]
    used = []

    def build(_model, _path, scale):
        used.append(scale)
        return 3

    monkeypatch.setattr(design, "_build_cv_pdf", build)
    monkeypatch.setattr(design, "_pdf_last_page_fill_ratio", lambda _: .1)
    result = design.render_cv_pdf({}, tmp_path / "unused.pdf")
    assert used == [design.CV_PDF_LAYOUTS[k] for k in ("comfortable", "compact", "comfortable")]
    assert result == design.CV_PDF_LAYOUTS["comfortable"]


def test_dashboard_routes_use_the_real_renderer_without_touching_canonical_data(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    from dashboard import server

    before = canonical_profile_path().read_bytes()
    job = vacancy(1)
    monkeypatch.setattr(server, "ROOT", tmp_path)
    monkeypatch.setattr(server, "get_job_by_id", lambda _: job)
    monkeypatch.setattr(server, "log_medical_match", lambda *_: None)
    monkeypatch.setattr(server, "update_tailored_resume", lambda *_: None)
    with TestClient(server.app) as client:
        response = client.post("/api/master-cv")
        assert response.status_code == 200
        master = response.json()
        assert master["no_submission_performed"]
        assert master["review_warnings"]  # held drafts remain owner-visible
        response = client.post(f"/api/jobs/{job['id']}/prepare")
        assert response.status_code == 200
        paths = response.json()["documents"]["tailored_cv"]
    for output in (master["documents"], paths):
        with pdfplumber.open(output["pdf"]) as pdf:
            assert_geometry(pdf)
    assert canonical_profile_path().read_bytes() == before


def test_real_sparse_page_is_removed_by_spacing_only_and_word_uses_same_scale(tmp_path):
    model = {
        "name": "Rendering Test", "headline": "Whitespace compaction fixture",
        "certifications": [f"Training {i:02d} — verified fixture text" for i in range(38)],
    }
    before = tmp_path / "comfortable.pdf"
    assert design._build_cv_pdf(model, before, design.CV_PDF_LAYOUTS["comfortable"]) == 2
    assert design._pdf_last_page_fill_ratio(before) < .1
    paths = design.render_professional_document_artifacts(
        "Rendering fixture", tmp_path / "compacted", document_type="cv", canonical_cv_model=model,
    )
    with pdfplumber.open(paths["pdf"]) as pdf:
        assert len(pdf.pages) == 1
        assert_geometry(pdf)
        for line in model["certifications"]:
            assert line in flat(pdf.pages[0].extract_text())
        training = next(line for line in pdf.pages[0].extract_text_lines() if "Training 00" in line["text"])
        assert {round(c["size"], 1) for c in training["chars"] if c["text"] != "•"} == {10.5}
    doc = Document(paths["docx"])
    assert doc.styles["CV Body"].font.size.pt == 10.5
    assert doc.styles["CV Bullet"].paragraph_format.line_spacing.pt == pytest.approx(13.6)
    assert doc.styles["CV Bullet"].paragraph_format.space_after.pt == 2
