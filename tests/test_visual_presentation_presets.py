"""Both employer-facing formats preserve content and owner-selected visual identity."""

import pdfplumber
from docx import Document

from utils.document_design import render_professional_document_artifacts


CV_MODEL = {
    "name": "Jane Doe",
    "headline": "Medical Doctor and Nutrition Specialist",
    "application_focus": "Provincial Medical Coordinator",
    "contact_lines": ["Kabul, Afghanistan", "jane@example.org"],
    "profile": "Qualified physician with documented experience in nutrition and primary care.",
    "experience": [
        {
            "role": "TFU Medical Doctor",
            "org": "Example NGO",
            "loc": "Bamyan",
            "dates": "2022–2024",
            "bullets": [
                "Delivered clinical assessment and follow-up for children under verified protocols.",
                "Maintained medical reporting and supported safeguarding activities.",
            ],
        },
    ],
    "expertise": [{"group": "Clinical Care", "items": ["IMAM", "HMIS", "Clinical reporting"]}],
    "education": ["Doctor of Medicine — Medical College — 2014–2020"],
    "certifications": ["Safeguarding & PSEA training"],
    "registration": ["Verified professional registration"],
    "exit_exam": ["Medical exit examination completed"],
    "languages": [("English", "Fluent"), ("Dari", "Native")],
}


def test_each_cv_palette_produces_searchable_text_and_matching_word_title(tmp_path):
    colors = {}
    for style in ("clinical", "executive"):
        model = {**CV_MODEL, "cv_style": style, "cv_layout": "comfortable"}
        result = render_professional_document_artifacts(
            "Jane Doe\nMEDICAL DOCTOR", tmp_path / style / "tailored_cv",
            document_type="cv", canonical_cv_model=model,
        )
        document = Document(result["docx"])
        with pdfplumber.open(result["pdf"]) as pdf:
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
            assert "TFU Medical Doctor" in text
            assert "Delivered clinical assessment" in text
            assert len(pdf.pages) >= 1
        assert any("TFU Medical Doctor" in paragraph.text for paragraph in document.paragraphs)
        colors[style] = str(document.styles["CV Professional Title"].font.color.rgb)
        assert all(result[key].endswith("." + key) for key in ("pdf", "docx", "txt"))
    assert colors["clinical"] != colors["executive"]


def test_cover_letter_and_cv_selected_theme_match(tmp_path):
    source = (
        "Subject: Provincial Medical Coordinator\n\n"
        "Dear Hiring Committee,\n\n"
        "I am applying to coordinate verified health programming.\n\n"
        "I have documented clinical assessment and monitoring experience.\n\n"
        "Sincerely,\nJane Doe\n"
    )
    result = render_professional_document_artifacts(
        source, tmp_path / "cover_letter", document_type="cover_letter",
        metadata={"tailored_documents": {"tailored_cv_model": {"cv_style": "executive"}}},
    )
    document = Document(result["docx"])
    assert str(document.styles["JF Title"].font.color.rgb).upper() == "8A632D"
    with pdfplumber.open(result["pdf"]) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
        assert "Jane Doe" in text
        assert "Provincial Medical Coordinator" in text
        assert "I am applying" in text
